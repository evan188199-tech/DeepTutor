"""Characterization tests for the silent error-swallowing branches in
``deeptutor/co_writer/docx_converter.py``.

DT-22 (AGEN-412) flagged three swallows with no behavior lock, where a
malformed document loses formatting or rows without any signal:

- ``_style_name``        — ``except Exception: pass`` → ``""``
- ``_run_to_markdown``   — ``run.font.strike`` read failures dropped silently
- ``_table_to_markdown`` — ``except Exception: continue`` → row dropped

Each test documents CURRENT behavior on the current baseline. Assertions
marked *CURRENT* lock today's silent-loss outcome and are expected to be
flipped by the corresponding fix (for tables, the fix shape is PR #1700:
warn + keep a placeholder row; for strike, emit ``~~...~~``). Tests marked
*control* prove the surrounding machinery works, so a fix card can trust
that a flipped assertion fails for the right reason.

Fixtures use real python-docx documents where a malformed document can
trigger the branch on the installed python-docx (1.2.0: invalid
``w:strike/@w:val`` raises ``InvalidXmlError``; orphaned style ids and
malformed table grids degrade upstream instead of raising). Branches that
python-docx 1.2.0 no longer reaches are locked with minimal stub objects,
which also pins the module contract for older python-docx versions where
those code paths did raise (e.g. KeyError for orphaned style ids).
"""

from io import BytesIO

from docx import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest

from deeptutor.co_writer.docx_converter import (
    _cell_text,
    _paragraph_to_markdown,
    _run_to_markdown,
    _style_name,
    _table_to_markdown,
    docx_to_markdown,
)

# --------------------------------------------------------------------------
# Stub helpers: minimal duck-typed paragraph / run / table objects used to
# reach ``except`` branches that the installed python-docx no longer raises
# into from real documents.
# --------------------------------------------------------------------------


class _ExplodingStyle:
    @property
    def name(self):
        raise RuntimeError("style.name lookup exploded")


class _ParagraphStyleRaises:
    text = "Heading text"
    runs: list = []

    @property
    def style(self):
        raise RuntimeError("style lookup exploded")

    def iter_inner_content(self):
        return iter([])


class _ParagraphStyleNameRaises:
    text = "Heading text"
    runs: list = []
    style = _ExplodingStyle()

    def iter_inner_content(self):
        return iter([])


class _ExplodingStrikeFont:
    @property
    def strike(self):
        raise RuntimeError("strike read exploded")


class _RunStrikeRaises:
    text = "gone"
    font = _ExplodingStrikeFont()
    bold = True
    italic = False


class _Cell:
    def __init__(self, text):
        self.text = text


class _Row:
    def __init__(self, texts):
        self.cells = [_Cell(t) for t in texts]


class _ExplodingRow:
    @property
    def cells(self):
        raise RuntimeError("row cells read exploded")


# --------------------------------------------------------------------------
# _style_name — swallow branch (returns "" on any style lookup failure)
# --------------------------------------------------------------------------


def test_style_name_returns_empty_when_style_property_raises():
    # CURRENT: any exception from the style lookup is swallowed to "".
    assert _style_name(_ParagraphStyleRaises()) == ""


def test_style_name_returns_empty_when_style_name_raises():
    # CURRENT: a failure while reading ``style.name`` is also swallowed.
    assert _style_name(_ParagraphStyleNameRaises()) == ""


def test_style_name_returns_empty_for_none_style():
    class _Para:
        style = None

    assert _style_name(_Para()) == ""


def test_style_name_returns_name_control():
    class _NamedStyle:
        name = "Heading 1"

    class _Para:
        style = _NamedStyle()

    assert _style_name(_Para()) == "Heading 1"


def test_paragraph_with_raising_style_renders_plain_body():
    # CURRENT: the heading level is silently lost — the body is still
    # converted, but no ``#`` markup and no signal that a style existed.
    out = _paragraph_to_markdown(_ParagraphStyleRaises(), {})
    assert out == "Heading text"
    assert not out.startswith("#")


# --------------------------------------------------------------------------
# _style_name — real-document end-to-end behavior
# --------------------------------------------------------------------------


def _orphan_style_docx() -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("orphan heading text", style="Heading 1")
    doc.add_paragraph("plain follower")
    styles_el = doc.styles.element
    for style_el in styles_el.findall(qn("w:style")):
        if style_el.get(qn("w:styleId")) == "Heading1":
            styles_el.remove(style_el)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_to_markdown_orphan_style_keeps_text_drops_heading():
    # Real document whose style definition is missing. On python-docx 1.2.0
    # the lookup degrades to the default style before our code runs, so the
    # loss path here is the silent "Normal" fallback rather than the
    # ``_style_name`` except branch. CURRENT on both routes: the text
    # survives, the heading markup does not, and nothing signals the loss.
    markdown = docx_to_markdown(_orphan_style_docx())
    assert "orphan heading text" in markdown
    assert "\n# " not in markdown
    assert not markdown.startswith("# ")


# --------------------------------------------------------------------------
# _run_to_markdown — strike branch (font.strike read failures dropped)
# --------------------------------------------------------------------------


def _docx_with_strike(val):
    doc = DocxDocument()
    run = doc.add_paragraph("struck text").runs[0]
    strike = OxmlElement("w:strike")
    if val is not None:
        strike.set(qn("w:val"), val)
    run._r.get_or_add_rPr().append(strike)
    buf = BytesIO()
    doc.save(buf)
    return DocxDocument(BytesIO(buf.getvalue())).paragraphs[0].runs[0]


def test_run_with_invalid_strike_value_drops_strike_markup():
    # Real malformed document: ``w:val="maybe"`` makes the strike read raise
    # InvalidXmlError. CURRENT: the failure is swallowed, the run converts
    # as plain text with no ``~~`` markers and no signal.
    run = _docx_with_strike("maybe")
    with pytest.raises(Exception):
        run.font.strike  # noqa: B018 — prove the fixture really is malformed
    assert _run_to_markdown(run) == "struck text"


def test_run_with_valid_strike_emits_strikethrough_control():
    # Control: with a well-formed value the strike markup is emitted, so a
    # flipped post-fix assertion cannot pass for the wrong reason.
    run = _docx_with_strike("true")
    assert _run_to_markdown(run) == "~~struck text~~"


def test_run_strike_read_failure_continues_with_bold():
    # CURRENT: after the swallowed strike read, conversion continues — the
    # bold read outside the try block still applies, without strike markers.
    assert _run_to_markdown(_RunStrikeRaises()) == "**gone**"


def test_docx_to_markdown_invalid_strike_keeps_text_without_markup():
    # End-to-end through the public entry: no exception escapes, output is
    # inspectable, and the strikethrough formatting silently disappeared.
    doc = DocxDocument()
    run = doc.add_paragraph("struck text").runs[0]
    strike = OxmlElement("w:strike")
    strike.set(qn("w:val"), "maybe")
    run._r.get_or_add_rPr().append(strike)
    buf = BytesIO()
    doc.save(buf)
    markdown = docx_to_markdown(buf.getvalue())
    assert "struck text" in markdown
    assert "~~" not in markdown


# --------------------------------------------------------------------------
# _table_to_markdown — bad-row branch (except Exception: continue)
# --------------------------------------------------------------------------


def test_table_bad_row_silently_dropped():
    # CURRENT: a row whose cells cannot be read is dropped without any
    # placeholder, warning, or marker — the markdown table just shrinks.
    table = type(
        "_Table",
        (),
        {"rows": [_Row(["a", "b"]), _ExplodingRow(), _Row(["c", "d"])]},
    )()
    assert _table_to_markdown(table) == "| a | b |\n| --- | --- |\n| c | d |"


def test_table_all_rows_bad_returns_empty():
    # CURRENT: if every row fails, the table converts to "" — total silent
    # loss; callers cannot distinguish an empty table from a broken one.
    table = type("_Table", (), {"rows": [_ExplodingRow(), _ExplodingRow()]})()
    assert _table_to_markdown(table) == ""


def test_table_deleted_tc_lenient_on_current_python_docx():
    # Real document with a row missing a ``w:tc``. python-docx 1.2.0 reads
    # the remaining cells without raising, so the historical IndexError
    # route into the ``except`` branch is not reachable this way anymore.
    # Kept as a positive control: the row survives, right-padded to the
    # grid width — and if a future python-docx makes this strict again,
    # this test documents the then-current silent-drop behavior.
    doc = DocxDocument()
    table = doc.add_table(rows=2, cols=3)
    for row in range(2):
        for col in range(3):
            table.cell(row, col).text = "abcdef"[row * 3 + col]
    tr = table.rows[1]._tr
    tr.remove(tr.findall(qn("w:tc"))[1])
    buf = BytesIO()
    doc.save(buf)
    reloaded = DocxDocument(BytesIO(buf.getvalue())).tables[0]
    cells = [cell.text for cell in reloaded.rows[1].cells]  # must not raise
    assert cells == ["d", "f"]
    assert _table_to_markdown(reloaded) == ("| a | b | c |\n| --- | --- | --- |\n| d | f |  |")


def test_table_cell_text_escapes_pipes_control():
    # Control for fixture soundness: cell text with pipes survives escaped.
    assert _cell_text(_Cell("a | b")) == r"a \| b"
    assert _cell_text(_Cell("line1\nline2")) == "line1 line2"


# --------------------------------------------------------------------------
# Cross-branch sanity through the public entry
# --------------------------------------------------------------------------


def test_docx_to_markdown_never_raises_on_malformed_formatting():
    # A document combining an invalid strike value still converts end to
    # end: no exception, output inspectable, formatting silently reduced.
    doc = DocxDocument()
    doc.add_paragraph("kept title text", style="Heading 1")
    run = doc.add_paragraph("struck text").runs[0]
    strike = OxmlElement("w:strike")
    strike.set(qn("w:val"), "maybe")
    run._r.get_or_add_rPr().append(strike)
    buf = BytesIO()
    doc.save(buf)
    markdown = docx_to_markdown(buf.getvalue())
    assert "kept title text" in markdown
    assert "struck text" in markdown
